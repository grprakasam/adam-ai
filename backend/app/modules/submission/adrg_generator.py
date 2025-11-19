"""
ADRG (Analysis Data Reviewer's Guide) generator
Creates DOCX document from template
"""

from typing import List, Optional
from pathlib import Path
from datetime import datetime
import hashlib
from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH

from .models import StudyInfo, DatasetInfo, DocumentResponse, ArtifactRef


class ADRGGenerator:
    """
    Generates ADRG (Analysis Data Reviewer's Guide) in DOCX format
    Provides placeholder template with study-specific information
    """
    
    def generate(
        self,
        study_info: StudyInfo,
        datasets: List[DatasetInfo],
        notes: Optional[str] = None,
        output_dir: Optional[Path] = None
    ) -> DocumentResponse:
        """
        Generate ADRG document
        
        Args:
            study_info: Study metadata
            datasets: List of dataset information
            notes: Optional additional notes
            output_dir: Output directory
        
        Returns:
            DocumentResponse with artifact reference
        """
        # Create output directory
        if output_dir is None:
            output_dir = Path(__file__).parents[4] / "Content" / "Submissions" / study_info.study_id / "docs"
        
        output_dir.mkdir(parents=True, exist_ok=True)
        
        # Create document
        doc = Document()
        
        # Add title
        title = doc.add_heading("Analysis Data Reviewer's Guide (ADRG)", level=1)
        title.alignment = WD_ALIGN_PARAGRAPH.CENTER
        
        # Study information
        doc.add_heading("1. Study Information", level=2)
        doc.add_paragraph(f"Study ID: {study_info.study_id}")
        doc.add_paragraph(f"Protocol Number: {study_info.protocol_number}")
        doc.add_paragraph(f"Study Name: {study_info.study_name}")
        doc.add_paragraph(f"Sponsor: {study_info.sponsor}")
        
        if study_info.description:
            doc.add_paragraph(f"Description: {study_info.description}")
        
        # Standards
        doc.add_heading("2. Standards and Versions", level=2)
        doc.add_paragraph(f"ADaM IG Version: {study_info.adam_ig_version}")
        if study_info.meddra_version:
            doc.add_paragraph(f"MedDRA Version: {study_info.meddra_version}")
        
        # Datasets overview
        doc.add_heading("3. Analysis Datasets", level=2)
        doc.add_paragraph(f"This submission includes {len(datasets)} ADAM dataset(s):")
        
        # Add table of datasets
        table = doc.add_table(rows=1, cols=3)
        table.style = 'Light Grid Accent 1'
        
        hdr_cells = table.rows[0].cells
        hdr_cells[0].text = 'Dataset'
        hdr_cells[1].text = 'Label'
        hdr_cells[2].text = 'Structure'
        
        for dataset in datasets:
            row_cells = table.add_row().cells
            row_cells[0].text = dataset.name
            row_cells[1].text = dataset.label
            row_cells[2].text = dataset.structure
        
        # Derivation summary
        doc.add_heading("4. Derivation Approach", level=2)
        doc.add_paragraph(
            "Analysis datasets were derived from SDTM datasets using validated programming methods. "
            "All derivations follow CDISC ADaM Implementation Guide principles."
        )
        
        if notes:
            doc.add_paragraph(f"\nAdditional Notes:\n{notes}")
        
        # Validation summary
        doc.add_heading("5. Quality Control", level=2)
        doc.add_paragraph(
            "All analysis datasets have been validated for CDISC compliance and data integrity. "
            "Validation was performed using Pinnacle 21 Community Edition and internal QC procedures."
        )
        
        # References
        doc.add_heading("6. References", level=2)
        doc.add_paragraph("• CDISC Analysis Data Model (ADaM) Implementation Guide")
        doc.add_paragraph("• Define-XML Specification 2.1")
        doc.add_paragraph("• Study Protocol and Statistical Analysis Plan")
        
        # Footer
        doc.add_paragraph(f"\n\nGenerated: {datetime.utcnow().strftime('%Y-%m-%d %H:%M:%S')} UTC")
        doc.add_paragraph("Note: This document is a template. Please review and customize as needed.")
        
        # Save document
        output_file = output_dir / f"adrg_{study_info.study_id}.docx"
        doc.save(str(output_file))
        
        # Calculate checksum
        with open(output_file, 'rb') as f:
            checksum = hashlib.sha256(f.read()).hexdigest()
        
        # Create artifact reference
        artifact = ArtifactRef(
            name=output_file.name,
            path=str(output_file),
            created_at=datetime.utcnow().isoformat(),
            checksum=checksum,
            size_bytes=output_file.stat().st_size
        )
        
        return DocumentResponse(artifact=artifact)
