"""
Text extraction from various document formats
Supports PDF, DOCX, XLSX, TXT
"""
import os
from pathlib import Path
from typing import Optional
from loguru import logger

try:
    import pdfplumber
    PDFPLUMBER_AVAILABLE = True
except ImportError:
    PDFPLUMBER_AVAILABLE = False
    logger.warning("pdfplumber not available, PDF extraction will fail")

try:
    from docx import Document
    DOCX_AVAILABLE = True
except ImportError:
    DOCX_AVAILABLE = False
    logger.warning("python-docx not available, DOCX extraction will fail")

try:
    import openpyxl
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False
    logger.warning("openpyxl not available, XLSX extraction will fail")


class TextExtractor:
    """Extract text from various document formats"""

    @staticmethod
    def extract_text(file_path: str, mime_type: Optional[str] = None) -> str:
        """
        Extract text from a document file

        Args:
            file_path: Path to the document file
            mime_type: MIME type of the file (optional)

        Returns:
            Extracted text content
        """
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        # Determine file type
        if mime_type:
            file_type = mime_type
        else:
            suffix = path.suffix.lower()
            mime_map = {
                ".pdf": "application/pdf",
                ".docx": "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                ".doc": "application/msword",
                ".xlsx": "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                ".xls": "application/vnd.ms-excel",
                ".txt": "text/plain"
            }
            file_type = mime_map.get(suffix, "text/plain")

        # Extract based on type
        if "pdf" in file_type:
            return TextExtractor._extract_pdf(file_path)
        elif "wordprocessingml" in file_type or "msword" in file_type:
            return TextExtractor._extract_docx(file_path)
        elif "spreadsheetml" in file_type or "ms-excel" in file_type:
            return TextExtractor._extract_xlsx(file_path)
        elif "text/plain" in file_type:
            return TextExtractor._extract_txt(file_path)
        else:
            logger.warning(f"Unsupported file type: {file_type}, attempting text extraction")
            return TextExtractor._extract_txt(file_path)

    @staticmethod
    def _extract_pdf(file_path: str) -> str:
        """Extract text from PDF"""
        if not PDFPLUMBER_AVAILABLE:
            raise ImportError("pdfplumber is required for PDF extraction")

        text_parts = []
        try:
            with pdfplumber.open(file_path) as pdf:
                for page in pdf.pages:
                    text = page.extract_text()
                    if text:
                        text_parts.append(text)
        except Exception as e:
            logger.error(f"Error extracting PDF: {e}")
            raise

        return "\n\n".join(text_parts)

    @staticmethod
    def _extract_docx(file_path: str) -> str:
        """Extract text from DOCX"""
        if not DOCX_AVAILABLE:
            raise ImportError("python-docx is required for DOCX extraction")

        try:
            doc = Document(file_path)
            paragraphs = [para.text for para in doc.paragraphs if para.text.strip()]
            return "\n".join(paragraphs)
        except Exception as e:
            logger.error(f"Error extracting DOCX: {e}")
            raise

    @staticmethod
    def _extract_xlsx(file_path: str) -> str:
        """Extract text from XLSX"""
        if not OPENPYXL_AVAILABLE:
            raise ImportError("openpyxl is required for XLSX extraction")

        text_parts = []
        try:
            workbook = openpyxl.load_workbook(file_path, data_only=True)
            for sheet_name in workbook.sheetnames:
                sheet = workbook[sheet_name]
                text_parts.append(f"Sheet: {sheet_name}")
                for row in sheet.iter_rows(values_only=True):
                    row_text = " | ".join(str(cell) if cell is not None else "" for cell in row)
                    if row_text.strip():
                        text_parts.append(row_text)
        except Exception as e:
            logger.error(f"Error extracting XLSX: {e}")
            raise

        return "\n".join(text_parts)

    @staticmethod
    def _extract_txt(file_path: str) -> str:
        """Extract text from plain text file"""
        try:
            with open(file_path, "r", encoding="utf-8", errors="ignore") as f:
                return f.read()
        except Exception as e:
            logger.error(f"Error extracting text: {e}")
            raise



